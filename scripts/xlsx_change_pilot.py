#!/usr/bin/env python3
"""Run a concurrent Change Pilot REST test from an XLSX change sheet.

The script reads raw change points, checks service health, then processes rows
with a configurable number of concurrent requests (``--concurrency``, default
5). Each row gets at most ``--retry`` retries after its initial request
(default 3). Credentials and provider payloads are never printed.
"""
from __future__ import annotations

import argparse
import concurrent.futures
import json
import re
import sys
import time
import urllib.error
import threading
import urllib.request
from dataclasses import dataclass, field as dataclass_field
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
WORKERS = 5
MAX_RETRIES = 3
REQUEST_TIMEOUT_SECONDS = 120
RETRY_BACKOFF_SECONDS = 0.5
CHANGE_POINT_HEADERS = ("原始变更点", "变更点", "变更内容")


@dataclass(frozen=True)
class ChangePoint:
    row_number: int
    text: str


class InputError(Exception):
    """Raised when the XLSX cannot provide usable change points."""


class RequestFailure(Exception):
    """Raised for a failed health or transformation request."""


@dataclass
class ApiStats:
    """Thread-safe counters for transformation API attempts."""

    total_calls: int = 0
    failed_calls: int = 0
    _lock: threading.Lock = dataclass_field(default_factory=threading.Lock, repr=False)

    def record(self, failed: bool) -> None:
        with self._lock:
            self.total_calls += 1
            if failed:
                self.failed_calls += 1


def load_token() -> str:
    env_path = SCRIPT_DIR / ".env"
    if not env_path.exists():
        return ""
    for line in env_path.read_text(encoding="utf-8").splitlines():
        if line.startswith("CHANGE_PILOT_API_TOKEN="):
            return line.split("=", 1)[1].strip()
    return ""


def _normalise(value: Any) -> str:
    return "" if value is None else "".join(str(value).split())


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_module_header(value: Any) -> bool:
    """True when *value* is a standalone module-column header.

    The normalised text must contain ``模块`` yet stay short (a header, not a
    long description such as ``安全模块固件版本号`` or ``各模块有单独版本号``).
    """
    normalised = _normalise(value)
    return "模块" in normalised and len(normalised) <= 6


def _module_column(sheet: Any, header_rows: int = 30) -> int | None:
    """Return the 1-based module column of *sheet*, else None if not found."""
    for row in sheet.iter_rows(max_row=header_rows):
        for cell in row:
            if cell.value is not None and _is_module_header(cell.value):
                return cell.column
    return None


# Category markers that head real change-point descriptions (as opposed to a
# short module name) in these change sheets, e.g. ``改进: 安全模块``,
# ``新功能: 触摸屏``, ``错误修复: 液晶``, ``辅助：编译相关``.
_CHANGE_POINT_PREFIX = re.compile(
    r"^\s*(改进|新功能|错误修复|修复|优化|新增|调整|补充|辅助|Bug|BUG|ID|版本|[#＃])",
    re.IGNORECASE,
)


def _looks_like_change_point(text: str) -> bool:
    """True when *text* reads as a change-point description rather than a short
    module name: it leads with a category marker or carries the ``详细信息``
    detail marker found in these sheets' change-point cells."""
    return bool(_CHANGE_POINT_PREFIX.match(text) or "详细信息" in text)


def read_change_points(path: Path) -> list[ChangePoint]:
    """Extract one raw change point per data row from an XLSX workbook."""
    if not path.exists():
        raise InputError(f"XLSX 文件不存在: {path}")
    if not path.is_file():
        raise InputError(f"XLSX 路径不是文件: {path}")
    if path.suffix.lower() != ".xlsx":
        raise InputError(f"仅支持 .xlsx 文件: {path}")

    try:
        workbook = load_workbook(path, read_only=False, data_only=True)
    except Exception as exc:  # noqa: BLE001 - normalize library/parser errors
        raise InputError(f"XLSX 文件无法读取，可能已加密或损坏: {path}") from exc

    try:
        sheets = list(workbook.worksheets)
        if not sheets:
            raise InputError("XLSX 中没有工作表")

        # Search every sheet for a change-point column and pick, across the whole
        # workbook, the header (row, column) whose cells read most like change
        # points. Ordinary (non-read-only) mode is used deliberately: some
        # workbooks carry an incorrect ``dimension`` so the read-only parser sees
        # only the first row, and merging sheets by row number also drops data.
        # Ranking by change-point-like content (not just row count) keeps a merged
        # section title such as ``变更点说明`` from out-ranking the real
        # ``变更点（有0A单、BUG编号的需注明）`` column, whose data rows can be
        # roughly as numerous as the module-name column's.
        best_score = (-1, -1)
        best_points: list[ChangePoint] = []
        for sheet in sheets:
            rows = list(sheet.iter_rows(values_only=True))
            for header_row, row in enumerate(rows):
                for header_index, value in enumerate(row):
                    if not any(name in _normalise(value) for name in CHANGE_POINT_HEADERS):
                        continue
                    points = []
                    for row_number, data in enumerate(rows[header_row + 1:], start=header_row + 2):
                        if header_index >= len(data):
                            continue
                        cell = _text(data[header_index])
                        if cell:
                            points.append(ChangePoint(row_number, cell))
                    if not points:
                        continue
                    score = (
                        sum(1 for point in points if _looks_like_change_point(point.text)),
                        len(points),
                    )
                    if score > best_score:
                        best_score = score
                        best_points = points
        if best_points:
            return best_points

        names = ", ".join(sheet.title for sheet in sheets)
        raise InputError(f"XLSX 中没有可用的原始变更点（工作表: {names}）")
    finally:
        workbook.close()


def write_refined_to_module_column(
    source_path: Path, results: list[tuple[ChangePoint, str | None]]
) -> Path | None:
    """Overwrite each row's module column with its refined change point.

    Every successful point replaces the module-column cell of the row it came
    from. A failed point leaves its module cell untouched. The result is saved
    as a new workbook named ``<source_stem>_AI.xlsx``; the source is never
    modified.

    Returns the output path, or ``None`` when no sheet exposes a module column
    or no point succeeded.
    """
    successful = [(point, output) for point, output in results if output]
    if not successful:
        return None
    try:
        workbook = load_workbook(source_path, data_only=False)
    except Exception as exc:  # noqa: BLE001 - normalize library/parser errors
        raise InputError(f"XLSX 无法以可写模式打开: {source_path}") from exc
    try:
        target: Any = None
        column: int | None = None
        for sheet in workbook.worksheets:
            found = _module_column(sheet)
            if found is not None:
                target, column = sheet, found
                break
        if target is None or column is None:
            return None
        for point, output in successful:
            target.cell(row=point.row_number, column=column).value = output
        output_path = source_path.with_name(f"{source_path.stem}_AI.xlsx")
        try:
            workbook.save(output_path)
        except OSError as exc:
            raise InputError(
                f"无法写入输出文件 {output_path}: {exc}。"
                "若该文件已在 Excel 中打开，请先关闭后再重试。"
            ) from exc
        return output_path
    finally:
        workbook.close()


def _request_json(
    url: str,
    token: str,
    method: str = "GET",
    payload: dict[str, Any] | None = None,
) -> dict[str, Any]:
    data = None
    headers = {"Authorization": f"Bearer {token}"}
    if payload is not None:
        data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        headers["Content-Type"] = "application/json"
    request = urllib.request.Request(url, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(request, timeout=REQUEST_TIMEOUT_SECONDS) as response:
            if not 200 <= response.status < 300:
                raise RequestFailure(f"HTTP {response.status}")
            body = response.read().decode("utf-8")
    except urllib.error.HTTPError as exc:
        raise RequestFailure(f"HTTP {exc.code}") from exc
    except (urllib.error.URLError, TimeoutError, OSError) as exc:
        raise RequestFailure(exc.__class__.__name__) from exc
    try:
        parsed = json.loads(body)
    except (TypeError, ValueError) as exc:
        raise RequestFailure("响应不是有效 JSON") from exc
    if not isinstance(parsed, dict):
        raise RequestFailure("响应 JSON 不是对象")
    return parsed


def health_check(base_url: str, token: str) -> str:
    try:
        result = _request_json(f"{base_url}/health", token)
        if result.get("status") != "ok":
            raise RequestFailure("health status 非 ok")
        provider = result.get("provider")
        model = provider.get("model") if isinstance(provider, dict) else None
        return model.strip() if isinstance(model, str) and model.strip() else "未配置"
    except RequestFailure as exc:
        raise RequestFailure(f"健康检查失败: {exc}") from exc


def call_change_pilot(base_url: str, token: str, raw_text: str) -> str:
    result = _request_json(
        f"{base_url}/v1/change-pilot",
        token,
        method="POST",
        payload={"raw_text": raw_text, "mode": "default"},
    )
    output = result.get("customer_output")
    if result.get("success") is not True or not isinstance(output, str) or not output.strip():
        raise RequestFailure("接口返回失败")
    return output.strip()


def process_one(
    base_url: str,
    token: str,
    point: ChangePoint,
    stats: ApiStats,
    max_retries: int = MAX_RETRIES,
    progress: "Callable[[int], None] | None" = None,
) -> tuple[ChangePoint, str | None]:
    """Transform one point, retrying up to ``max_retries`` extra times.

    ``progress`` is an optional callback invoked as ``progress(attempt)``
    before each retry (attempt = 1-based retry count) so the caller can
    surface the retry in its progress view.
    """
    for attempt in range(max_retries + 1):
        if attempt and progress is not None:
            progress(attempt)
        try:
            output = call_change_pilot(base_url, token, point.text)
            stats.record(failed=False)
            return point, output
        except RequestFailure:
            stats.record(failed=True)
            if attempt >= max_retries:
                break
            time.sleep(RETRY_BACKOFF_SECONDS * (attempt + 1))
    return point, None


def process_in_batches(
    base_url: str,
    token: str,
    points: list[ChangePoint],
    stats: ApiStats,
    workers: int = WORKERS,
    max_retries: int = MAX_RETRIES,
) -> list[tuple[ChangePoint, str | None]]:
    """Keep the configured number of requests in flight and refresh progress slots."""
    results: list[tuple[ChangePoint, str | None] | None] = [None] * len(points)
    slots: list[str] = []
    slot_futures: dict[concurrent.futures.Future, tuple[int, int]] = {}
    next_index = 0
    interactive = sys.stdout.isatty()
    block_rendered = False
    terminal_lock = threading.Lock()

    def paint() -> None:
        # Assumes terminal_lock is held. Rise back to the block's first line
        # only after it has been drawn once: on the first render the cursor is
        # already at the block start (just below the health-check log), so
        # rising early would climb into those logs and wipe them.
        nonlocal block_rendered
        if interactive:
            if block_rendered:
                sys.stdout.write(f"\033[{len(slots)}A")
            else:
                block_rendered = True
            for message in slots:
                sys.stdout.write(f"\033[2K{message}\n")
        else:
            for message in slots:
                print(message)
        sys.stdout.flush()

    def render() -> None:
        with terminal_lock:
            paint()

    def report_retry(slot: int, ordinal: int, attempt: int) -> None:
        """Flip a slot to its retry message and repaint immediately."""
        with terminal_lock:
            slots[slot] = f"正在处理第{ordinal}条数据，第{attempt}次重试中..."
            paint()

    def make_progress(slot: int, ordinal: int) -> "Callable[[int], None]":
        return lambda attempt: report_retry(slot, ordinal, attempt)

    def submit_one(point_index: int, slot: int) -> concurrent.futures.Future:
        point = points[point_index]
        return executor.submit(
            process_one, base_url, token, point, stats, max_retries,
            make_progress(slot, point_index + 1),
        )

    with concurrent.futures.ThreadPoolExecutor(max_workers=workers) as executor:
        while next_index < len(points) and len(slot_futures) < workers:
            slot = len(slot_futures)
            slots.append(f"正在处理第{next_index + 1}条数据...")
            slot_futures[submit_one(next_index, slot)] = (slot, next_index)
            next_index += 1
        render()

        while slot_futures:
            done, _ = concurrent.futures.wait(
                slot_futures, return_when=concurrent.futures.FIRST_COMPLETED
            )
            for future in done:
                slot, point_index = slot_futures.pop(future)
                point, output = future.result()
                results[point_index] = (point, output)
                with terminal_lock:
                    if next_index < len(points):
                        slots[slot] = f"正在处理第{next_index + 1}条数据..."
                        slot_futures[submit_one(next_index, slot)] = (slot, next_index)
                        next_index += 1
                    else:
                        slots[slot] = "处理完成，等待其他任务..."
            render()

    if interactive and slots:
        with terminal_lock:
            sys.stdout.write(f"\033[{len(slots)}A")
            for _ in slots:
                sys.stdout.write("\033[2K\n")
            sys.stdout.write("处理进度完成\n")
            sys.stdout.flush()
    return [result for result in results if result is not None]


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "从 XLSX 变更单读取原始变更点，检查服务健康状态后，以固定并发请求数"
            "批量调用 Change Pilot 接口提炼客户可见的变更点。"
        ),
        epilog=(
            "示例：\n"
            "  python3 scripts/xlsx_change_pilot.py --xlsx 变更单.xlsx "
            "--base-url http://127.0.0.1:18081\n"
            "  python3 scripts/xlsx_change_pilot.py --xlsx 变更单.xlsx "
            "--base-url http://127.0.0.1:18081 --concurrency 10 --retry 2\n\n"
            "接口令牌从脚本同目录 .env 的 CHANGE_PILOT_API_TOKEN 读取，不会打印。"
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument(
        "--xlsx", type=Path, required=True,
        help="原始变更单 .xlsx 路径",
    )
    parser.add_argument(
        "--base-url", required=True,
        help="Change Pilot 服务地址，例如 http://127.0.0.1:18081",
    )
    parser.add_argument(
        "--concurrency", type=int, default=WORKERS,
        help=f"并发请求数（默认 {WORKERS}）",
    )
    parser.add_argument(
        "--retry", type=int, default=MAX_RETRIES,
        help=f"每条变更点最大重试次数，0 表示不重试（默认 {MAX_RETRIES}）",
    )
    args = parser.parse_args(argv)
    if args.concurrency < 1:
        parser.error(f"--concurrency 必须 >= 1，received {args.concurrency}")
    if args.retry < 0:
        parser.error(f"--retry 必须 >= 0，received {args.retry}")
    base_url = args.base_url.rstrip("/")

    try:
        points = read_change_points(args.xlsx)
    except InputError as exc:
        print(f"文件校验失败: {exc}", file=sys.stderr)
        return 2
    print(f"文件校验通过: {args.xlsx}，读取到 {len(points)} 条原始变更点")

    token = load_token()
    if not token:
        print("服务检查失败: .env 中未找到 CHANGE_PILOT_API_TOKEN", file=sys.stderr)
        return 2
    try:
        model = health_check(base_url, token)
    except RequestFailure as exc:
        print(str(exc), file=sys.stderr)
        return 2
    print(f"健康检查通过，当前模型: {model}")

    stats = ApiStats()
    started_at = time.perf_counter()
    results = process_in_batches(
        base_url, token, points, stats,
        workers=args.concurrency, max_retries=args.retry,
    )
    elapsed_seconds = time.perf_counter() - started_at
    success_count = 0
    for serial, (point, output) in enumerate(results, start=1):
        if output is None:
            print(f"{serial}. 大模型提炼失败")
        else:
            success_count += 1
            print(f"{serial}. {output}")
    print(f"成功数/总数: {success_count}/{len(results)}")
    print(f"总耗时: {elapsed_seconds:.2f} 秒")
    print(f"API 总调用次数: {stats.total_calls}")
    print(f"API 失败次数: {stats.failed_calls}")
    try:
        output_path = write_refined_to_module_column(args.xlsx, results)
    except InputError as exc:
        print(str(exc), file=sys.stderr)
        return 2
    if output_path is None:
        print("提醒: 未找到含'模块'列头的工作表，未写回结果（仅 console 输出）")
    else:
        print(f"已写回模块列，输出文件: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
